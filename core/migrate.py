"""Explicit, verified workbook import. Sources are always left unchanged.

Supports this app's monthly sheets and the earlier single Timesheet format.
Missing dates require user confirmation. Imports are written to a temporary
file, verified, then installed after backing up the source and local data.
No workbook is discovered, renamed or imported automatically.
"""

import collections
import datetime as dt
import math
import os
import shutil
import uuid
from dataclasses import dataclass, field
from typing import Optional

from openpyxl import load_workbook

from . import config as cfgmod
from .excel_store import Entry, ExcelStore, HEADERS, read_data_rows
from .timeutil import month_key, month_display


@dataclass
class DatelessRow:
    row: int
    day_name: str
    start: Optional[dt.time]
    finish: Optional[dt.time]
    break_min: Optional[int]
    overtime: float
    type: str
    summary: str
    suggested: Optional[dt.date] = None


@dataclass
class MigrationPlan:
    entries: list = field(default_factory=list)
    dateless: list = field(default_factory=list)


def _check_headers(ws):
    headers = [ws.cell(row=1, column=i).value for i in range(1, 12)]
    if headers != HEADERS:
        raise ValueError(f"Sheet '{ws.title}' does not have the supported tracker columns.")


def _check_raw_rows(ws, legacy=False):
    """Reject ambiguous input instead of silently converting it to blank fields."""
    weekdays = {"Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"}
    for row in range(2, ws.max_row + 1):
        date = ws.cell(row=row, column=1).value
        day = ws.cell(row=row, column=2).value
        if date == "Totals" and day in (None, ""):
            continue
        if not isinstance(date, (dt.date, dt.datetime)) and day not in weekdays:
            if any(ws.cell(row=row, column=column).value not in (None, "")
                   for column in (3, 4, 5, 7, 8, 11)):
                raise ValueError(f"Sheet '{ws.title}', row {row}: enter a valid date before importing.")
            continue
        if not legacy and not isinstance(date, (dt.date, dt.datetime)):
            raise ValueError(f"Sheet '{ws.title}', row {row}: monthly sheets require dates on every entry.")
        for column in (3, 4):
            time = ws.cell(row=row, column=column).value
            if time not in (None, "") and not isinstance(time, (dt.time, dt.datetime)):
                raise ValueError(f"Sheet '{ws.title}', row {row}: start/finish must be Excel time cells or blank.")
        try:
            break_value = ws.cell(row=row, column=5).value
            if break_value not in (None, ""):
                break_number = float(break_value)
                if not math.isfinite(break_number) or break_number < 0 or not break_number.is_integer():
                    raise ValueError()
            overtime = float(ws.cell(row=row, column=7).value or 0)
            if not math.isfinite(overtime) or overtime < 0:
                raise ValueError()
        except (ValueError, TypeError):
            raise ValueError(f"Sheet '{ws.title}', row {row}: invalid break or overtime value.") from None


def read_original(path):
    """Read only a user-selected tracker; refuse unrecognised workbook layouts."""
    wb = load_workbook(path)
    plan = MigrationPlan()
    try:
        if "Timesheet" not in wb.sheetnames:
            sheets = [wb[name] for name in wb.sheetnames if name != "Summary"]
            if not sheets:
                raise ValueError("This workbook has no timesheet data sheets.")
            for ws in sheets:
                _check_headers(ws)
                _check_raw_rows(ws)
                plan.entries.extend(read_data_rows(ws))
        else:
            ws = wb["Timesheet"]
            _check_headers(ws)
            _check_raw_rows(ws, legacy=True)
            weekdays = {"Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
                        "Saturday", "Sunday"}
            for row in range(2, ws.max_row + 1):
                d = ws.cell(row=row, column=1).value
                if isinstance(d, dt.datetime):
                    d = d.date()
                day_name = ws.cell(row=row, column=2).value
                if not isinstance(d, dt.date) and day_name not in weekdays:
                    continue
                values = dict(
                    start=_t(ws.cell(row=row, column=3).value),
                    finish=_t(ws.cell(row=row, column=4).value),
                    break_min=_break(ws.cell(row=row, column=5).value),
                    overtime=float(ws.cell(row=row, column=7).value or 0),
                    type=str(ws.cell(row=row, column=8).value or "Office"),
                    summary=str(ws.cell(row=row, column=11).value or ""),
                )
                if isinstance(d, dt.date):
                    plan.entries.append(Entry(date=d, **values))
                else:
                    # Gaps elsewhere do not establish the date of this row.
                    plan.dateless.append(DatelessRow(row=row, day_name=str(day_name), **values))
    finally:
        wb.close()
    if not plan.entries and not plan.dateless:
        raise ValueError("There are no work entries in the selected workbook.")
    return plan


def _break(value):
    return int(value) if value not in (None, "") else None


def _t(value):
    if isinstance(value, dt.time):
        return value
    if isinstance(value, dt.datetime):
        return value.time()
    return None


def month_sums(entries):
    """{(year, month): (hours, overtime, total)} without Excel caches."""
    out = {}
    for entry in entries:
        key = month_key(entry.date)
        hours, overtime, total = out.get(key, (0.0, 0.0, 0.0))
        out[key] = (round(hours + entry.hours_worked(), 4),
                    round(overtime + (entry.overtime or 0), 4),
                    round(total + entry.total_hours(), 4))
    return out


def _entry_values(entry):
    return (entry.date, entry.start, entry.finish, entry.break_min,
            entry.overtime, entry.type, entry.summary)


def verify(original_entries, migrated_path, cfg):
    """Check every raw field and independently recomputed monthly hours/pay."""
    store = ExcelStore(path=migrated_path, cfg=cfg)
    store.load()
    migrated = store.entries
    original = sorted(original_entries, key=lambda entry: entry.date)
    report = []
    ok = True
    if [_entry_values(entry) for entry in original] != [
            _entry_values(entry) for entry in migrated]:
        ok = False
        report.append("FAIL: row counts or raw entry values differ from the source.")
    else:
        report.append(f"OK: all {len(original)} entry values preserved.")
    source_sums, migrated_sums = month_sums(original), month_sums(migrated)
    wb = load_workbook(migrated_path)
    try:
        for ws in wb.worksheets:
            if ws.title == "Summary":
                continue
            for row in range(2, ws.max_row + 1):
                date = ws.cell(row=row, column=1).value
                if isinstance(date, dt.datetime):
                    date = date.date()
                if not isinstance(date, dt.date):
                    continue
                try:
                    rate = cfgmod.rate_for_date(cfg, date)
                    expected = f"=I{row}*{cfgmod.format_rate(rate)}"
                except ValueError:
                    expected = None
                if ws.cell(row=row, column=10).value != expected:
                    ok = False
                    report.append(f"FAIL {date}: pay formula does not match the configured rate.")
    finally:
        wb.close()
    for key in sorted(set(source_sums) | set(migrated_sums)):
        if source_sums.get(key) != migrated_sums.get(key):
            ok = False
            report.append(f"FAIL {month_display(*key)}: computed hours differ.")
            continue
        report.append(f"OK {month_display(*key)}: hours, overtime and total agree.")
        try:
            source_pay = round(sum(entry.pay(cfgmod.rate_for_date(cfg, entry.date))
                                   for entry in original if month_key(entry.date) == key), 2)
            migrated_pay = round(sum(entry.pay(cfgmod.rate_for_date(cfg, entry.date))
                                     for entry in migrated if month_key(entry.date) == key), 2)
        except ValueError:
            report.append("Pay not calculated for this month: set effective rates in Settings.")
        else:
            if source_pay != migrated_pay:
                ok = False
                report.append(f"FAIL {month_display(*key)}: computed pay differs.")
            else:
                report.append(f"OK {month_display(*key)}: computed pay agrees.")
    return ok, report


def run(cfg, resolved, original_path, target_path=None):
    """Import a selected source, keeping source and target backups.

    ``resolved`` maps dateless row numbers to user-confirmed dates.
    Local data is replaced, never merged; the UI must confirm replacement.
    The source path is mandatory so no filesystem discovery can occur.
    """
    original_path = os.path.abspath(original_path)
    target_path = os.path.abspath(target_path or cfgmod.WORKBOOK_PATH)
    if os.path.normcase(os.path.realpath(original_path)) == os.path.normcase(
            os.path.realpath(target_path)):
        raise ValueError("Choose a different source; this is already the local workbook.")
    plan = read_original(original_path)
    unresolved = [row for row in plan.dateless if row.row not in resolved]
    if unresolved:
        return False, [f"FAIL: row {row.row} ({row.day_name}) needs a confirmed date."
                       for row in unresolved]
    entries = list(plan.entries)
    for row in plan.dateless:
        confirmed = resolved[row.row]
        if not isinstance(confirmed, dt.date) or isinstance(confirmed, dt.datetime):
            raise ValueError(f"Row {row.row} needs a date value.")
        entries.append(Entry(confirmed, row.start, row.finish, row.break_min,
                             row.overtime, row.type, row.summary))
    entries.sort(key=lambda entry: entry.date)
    duplicates = [date for date, count in collections.Counter(
        entry.date for entry in entries).items() if count > 1]
    if duplicates:
        return False, ["FAIL: duplicate dates: " + ", ".join(map(str, sorted(duplicates)))]

    cfgmod.ensure_dirs()
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    source_backup = os.path.join(cfgmod.BACKUP_DIR, f"Import_Source_{stamp}.xlsx")
    shutil.copy2(original_path, source_backup)
    temporary = target_path + f".{uuid.uuid4().hex}.import.xlsx"
    try:
        from .excel_store import build_workbook
        wb = build_workbook(entries, cfg)
        try:
            wb.save(temporary)
        finally:
            wb.close()
        ok, report = verify(entries, temporary, cfg)
        report.insert(0, f"Source backup: {source_backup}")
        if not ok:
            report.append("Import cancelled. Source and local workbook are unchanged.")
            return False, report
        if os.path.exists(target_path):
            old_backup = os.path.join(cfgmod.BACKUP_DIR, f"Before_Import_{stamp}.xlsx")
            shutil.copy2(target_path, old_backup)
            report.append(f"Previous local workbook backup: {old_backup}")
        os.replace(temporary, target_path)
        report.append("Import complete. The source workbook is unchanged.")
        return True, report
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
