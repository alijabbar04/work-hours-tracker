"""Workbook storage: one sheet per month + master Summary sheet.

The workbook is fully derived from the in-memory entry list, so every save
rebuilds it deterministically. Every write is preceded by a timestamped
backup (newest 30 kept) and followed by a one-way copy to the OneDrive
mirror folder (non-blocking on failure).
"""

import datetime as dt
import glob
import os
import shutil
import tempfile
from dataclasses import dataclass, field
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

from . import config as cfgmod
from .timeutil import (hours_between, month_key, month_sheet_name,
                       month_display, tax_year_of, tax_year_label)

# Excel styling constants (mirrors the original tracker exactly)
X_DARK, X_MINT, X_MINT_SOFT, X_GREY = "0D0D0D", "3DDC97", "E7F9F1", "F2F2F2"
X_WHITE, X_TXT, X_FONT = "FFFFFF", "1A1A1A", "Aptos"
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADERS = ["Date", "Day", "Start", "Finish", "Break (min)", "Hours worked",
           "Overtime (hrs)", "Type", "Total hours", "Pay", "Summary"]
COL_WIDTHS = [13, 11, 8, 8, 11, 13, 14, 18, 11, 11, 70]


@dataclass
class Entry:
    date: dt.date
    start: Optional[dt.time] = None
    finish: Optional[dt.time] = None
    break_min: Optional[int] = None
    overtime: float = 0.0
    type: str = "Office"
    summary: str = ""

    def hours_worked(self):
        return hours_between(self.start, self.finish, self.break_min)

    def total_hours(self):
        return round(self.hours_worked() + (self.overtime or 0), 4)

    def pay(self, rate):
        return round(self.total_hours() * rate, 2)


class ExcelStore:
    def __init__(self, path=None, cfg=None):
        self.path = path or cfgmod.WORKBOOK_PATH
        self.cfg = cfg if cfg is not None else cfgmod.load_config()
        self.entries = []  # list[Entry], kept sorted by date
        self.last_mirror_error = None

    # ----------------------------- reading -----------------------------
    def load(self):
        self.entries = []
        if not os.path.exists(self.path):
            return
        wb = load_workbook(self.path)
        for name in wb.sheetnames:
            if name == "Summary":
                continue
            self.entries.extend(read_data_rows(wb[name]))
        wb.close()
        self.entries.sort(key=lambda e: e.date)

    def months(self):
        """Sorted list of (year, month) present in the data."""
        return sorted({month_key(e.date) for e in self.entries})

    def entries_for_month(self, year, month):
        return [e for e in self.entries if month_key(e.date) == (year, month)]

    def entry_for_date(self, d):
        for e in self.entries:
            if e.date == d:
                return e
        return None

    # ----------------------------- writing -----------------------------
    def upsert(self, entry):
        previous = list(self.entries)
        self.entries = [e for e in self.entries if e.date != entry.date]
        self.entries.append(entry)
        self.entries.sort(key=lambda e: e.date)
        try:
            self.save()
        except Exception:
            self.entries = previous
            raise

    def delete(self, d):
        previous = list(self.entries)
        self.entries = [e for e in self.entries if e.date != d]
        try:
            self.save()
        except Exception:
            self.entries = previous
            raise

    def save(self):
        cfgmod.ensure_dirs()
        self.backup()
        wb = build_workbook(self.entries, self.cfg)
        # Write beside the destination and replace only after a complete save.
        # A failed export must never truncate the user's authoritative workbook.
        os.makedirs(os.path.dirname(os.path.abspath(self.path)), exist_ok=True)
        fd, temp_path = tempfile.mkstemp(prefix=".work-hours-", suffix=".xlsx",
                                         dir=os.path.dirname(os.path.abspath(self.path)))
        os.close(fd)
        try:
            wb.save(temp_path)
            os.replace(temp_path, self.path)  # PermissionError propagates to UI
        finally:
            wb.close()
            if os.path.exists(temp_path):
                os.remove(temp_path)
        self.mirror()
        self.export_mobile()

    def export_mobile(self):
        """Regenerate the phone companion app (never blocks a save)."""
        try:
            from . import mobile
            mobile.export_all(self.entries, self.cfg)
        except Exception:
            pass

    def backup(self):
        if not os.path.exists(self.path):
            return
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = os.path.join(cfgmod.BACKUP_DIR, f"Work_Hours_Tracker_{stamp}.xlsx")
        shutil.copy2(self.path, dest)
        prune_backups()

    def reconcile_mirror(self):
        """Catch-up sync at startup: if a previous save couldn't reach the
        OneDrive folder (offline, file locked), bring the mirror and the
        mobile app up to date now. The local workbook on C: is always
        authoritative - this never reads from OneDrive."""
        mirror_dir = self.cfg.get("onedrive_mirror_dir")
        if not (mirror_dir and os.path.exists(self.path)):
            return False
        dest = os.path.join(mirror_dir, os.path.basename(self.path))
        try:
            stale = (not os.path.exists(dest)
                     or os.path.getmtime(dest) + 1 < os.path.getmtime(self.path))
        except OSError:
            stale = True
        if not stale:
            return False
        ok = self.mirror()
        self.export_mobile()
        return ok

    def mirror(self):
        """One-way copy local -> OneDrive. Never blocks a save."""
        self.last_mirror_error = None
        mirror_dir = self.cfg.get("onedrive_mirror_dir")
        if not mirror_dir:
            return False
        try:
            os.makedirs(mirror_dir, exist_ok=True)
            shutil.copy2(self.path, os.path.join(mirror_dir, os.path.basename(self.path)))
            return True
        except OSError as e:
            self.last_mirror_error = str(e)
            return False


def prune_backups(pattern="Work_Hours_Tracker_*.xlsx", keep=cfgmod.BACKUP_KEEP):
    files = sorted(glob.glob(os.path.join(cfgmod.BACKUP_DIR, pattern)))
    for f in files[:-keep] if len(files) > keep else []:
        try:
            os.remove(f)
        except OSError:
            pass


# ----------------------------- row parsing -----------------------------
def read_data_rows(ws):
    """Parse Entry objects from a month sheet (raw values only, formulas ignored)."""
    out = []
    for row in range(2, ws.max_row + 1):
        d = ws.cell(row=row, column=1).value
        if isinstance(d, dt.datetime):
            d = d.date()
        if not isinstance(d, dt.date):
            continue  # blank spacer / footer rows
        start = _as_time(ws.cell(row=row, column=3).value)
        finish = _as_time(ws.cell(row=row, column=4).value)
        brk = ws.cell(row=row, column=5).value
        ot = ws.cell(row=row, column=7).value or 0
        typ = ws.cell(row=row, column=8).value or "Office"
        summary = ws.cell(row=row, column=11).value or ""
        out.append(Entry(date=d, start=start, finish=finish,
                         break_min=int(brk) if brk not in (None, "") else None,
                         overtime=float(ot), type=str(typ), summary=str(summary)))
    return out


def _as_time(v):
    if isinstance(v, dt.time):
        return v
    if isinstance(v, dt.datetime):
        return v.time()
    return None


# ----------------------------- workbook build -----------------------------
def build_workbook(entries, cfg):
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"

    months = sorted({month_key(e.date) for e in entries})
    month_ranges = {}  # (y, m) -> (sheet_name, first_data_row, last_data_row)
    for (y, m) in months:
        month_entries = sorted((e for e in entries if month_key(e.date) == (y, m)),
                               key=lambda e: e.date)
        name = month_sheet_name(y, m)
        ws = wb.create_sheet(name)
        last = write_month_sheet(ws, month_entries, cfg)
        month_ranges[(y, m)] = (name, 2, last)

    write_summary_sheet(summary_ws, months, month_ranges)
    return wb


def write_month_sheet(ws, month_entries, cfg):
    """Write header, data rows, blank row and totals footer. Returns last data row."""
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=X_FONT, bold=True, color=X_WHITE, size=11)
        cell.fill = PatternFill("solid", fgColor=X_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 26

    row = 1
    for e in month_entries:
        row += 1
        write_entry_row(ws, row, e, cfg)
    last_data = row

    footer = last_data + 2  # blank row between data and totals
    ws.cell(row=footer, column=1, value="Totals")
    ws.cell(row=footer, column=6, value=f"=SUM(F2:F{last_data})")
    ws.cell(row=footer, column=7, value=f"=SUM(G2:G{last_data})")
    ws.cell(row=footer, column=9, value=f"=SUM(I2:I{last_data})")
    ws.cell(row=footer, column=10,
            value=f'=IF(COUNT(J2:J{last_data})=ROWS(J2:J{last_data}),SUM(J2:J{last_data}),"")')
    for c in range(1, 12):
        cell = ws.cell(row=footer, column=c)
        cell.font = Font(name=X_FONT, bold=True, size=10, color=X_TXT)
        cell.fill = PatternFill("solid", fgColor=X_MINT_SOFT)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center")
        if c in (6, 7, 9):
            cell.number_format = "0.0"
        elif c == 10:
            cell.number_format = "#,##0.00"
    return last_data


def write_entry_row(ws, row, e, cfg):
    try:
        rate = cfgmod.rate_for_date(cfg, e.date)
    except ValueError:
        rate = None
    ws.cell(row=row, column=1, value=e.date).number_format = "ddd dd mmm"
    ws.cell(row=row, column=2, value=e.date.strftime("%A"))
    ws.cell(row=row, column=3, value=e.start).number_format = "HH:MM"
    ws.cell(row=row, column=4, value=e.finish).number_format = "HH:MM"
    ws.cell(row=row, column=5, value=e.break_min)
    ws.cell(row=row, column=6,
            value=f'=IF(OR(C{row}="",D{row}=""),0,(D{row}-C{row})*24-N(E{row})/60)'
            ).number_format = "0.0"
    ws.cell(row=row, column=7, value=e.overtime).number_format = "0.0"
    _write_text(ws.cell(row=row, column=8), e.type)
    ws.cell(row=row, column=9, value=f"=F{row}+G{row}").number_format = "0.0"
    pay_cell = ws.cell(row=row, column=10)
    if rate is not None:
        pay_cell.value = f"=I{row}*{cfgmod.format_rate(rate)}"
    else:
        pay_cell.comment = Comment("No rate configured for this date. Add one in Settings.",
                                   "Work Hours Tracker")
    pay_cell.number_format = "#,##0.00"
    _write_text(ws.cell(row=row, column=11), e.summary)

    weekend = "Weekend" in (e.type or "")
    n_bullets = max(1, (e.summary or "").count("•"))
    for c in range(1, 12):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=X_FONT, size=10, color=X_TXT)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(c == 11),
                                   horizontal="left" if c in (2, 8, 11) else "center")
        if weekend:
            cell.fill = PatternFill("solid", fgColor=X_MINT_SOFT)
        elif row % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=X_GREY)
        else:
            cell.fill = PatternFill(fill_type=None)
    ws.row_dimensions[row].height = max(30, 15 * n_bullets + 6)


def _write_text(cell, value):
    """Keep user text literal, including values beginning with '='.

    Explicit OOXML string cells prevent Excel from interpreting notes/types as
    formulas while preserving their text through a workbook round trip.
    """
    cell.value = str(value or "")
    cell.data_type = "s"


def write_summary_sheet(ws, months, month_ranges):
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEF", [22, 13, 14, 14, 13, 14]):
        ws.column_dimensions[col].width = w

    ws["A1"] = "Work Hours — Summary"
    ws["A1"].font = Font(name=X_FONT, bold=True, size=15, color=X_DARK)

    headers = ["Month", "Days logged", "Hours worked", "Overtime (hrs)",
               "Total hours", "Pay (£)"]
    hdr_row = 3
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=hdr_row, column=c, value=h)
        cell.font = Font(name=X_FONT, bold=True, color=X_WHITE, size=11)
        cell.fill = PatternFill("solid", fgColor=X_DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[hdr_row].height = 24

    row = hdr_row
    month_rows = {}
    for (y, m) in months:
        row += 1
        name, first, last = month_ranges[(y, m)]
        ref = f"'{name}'!"
        ws.cell(row=row, column=1, value=month_display(y, m))
        ws.cell(row=row, column=2, value=f"=COUNT({ref}A{first}:A{last})")
        ws.cell(row=row, column=3, value=f"=SUM({ref}F{first}:F{last})")
        ws.cell(row=row, column=4, value=f"=SUM({ref}G{first}:G{last})")
        ws.cell(row=row, column=5, value=f"=SUM({ref}I{first}:I{last})")
        ws.cell(row=row, column=6,
                value=f'=IF(COUNT({ref}J{first}:J{last})=ROWS({ref}J{first}:J{last}),SUM({ref}J{first}:J{last}),"")')
        _style_summary_row(ws, row, bold=False, banded=(row % 2 == 0))
        month_rows[(y, m)] = row
    first_month_row, last_month_row = hdr_row + 1, row

    if months:
        row += 1
        ws.cell(row=row, column=1, value="Grand total")
        for c in range(2, 7):
            col = get_column_letter(c)
            ws.cell(row=row, column=c,
                    value=(f'=IF(COUNT({col}{first_month_row}:{col}{last_month_row})='
                           f'ROWS({col}{first_month_row}:{col}{last_month_row}),'
                           f'SUM({col}{first_month_row}:{col}{last_month_row}),"")'
                           if c == 6 else
                           f"=SUM({col}{first_month_row}:{col}{last_month_row})"))
        _style_summary_row(ws, row, bold=True, fill=X_MINT_SOFT)

    # ---- tax-year section (UK: 6 Apr - 5 Apr) ----
    row += 2
    cell = ws.cell(row=row, column=1, value="Tax years (6 Apr – 5 Apr)")
    cell.font = Font(name=X_FONT, bold=True, size=12, color=X_DARK)
    row += 1
    for c, h in enumerate(["Tax year", "", "", "", "Total hours", "Gross pay (£)"], 1):
        cellh = ws.cell(row=row, column=c, value=h or None)
        cellh.font = Font(name=X_FONT, bold=True, color=X_WHITE, size=11)
        cellh.fill = PatternFill("solid", fgColor=X_DARK)
        cellh.border = BORDER
        cellh.alignment = Alignment(horizontal="center", vertical="center")

    # April spans two UK tax years. Derive years from actual entry dates and
    # aggregate rows by date criteria rather than assigning an entire month.
    tax_years = sorted({tax_year_of(wb_date)
                        for name, first, last in month_ranges.values()
                        for wb_row in range(first, last + 1)
                        if isinstance((wb_date := ws.parent[name].cell(wb_row, 1).value), dt.date)})
    for ty in tax_years:
        row += 1
        hour_sums, pay_sums, missing_rates = [], [], []
        for name, first, last in month_ranges.values():
            ref = f"'{name}'!"
            dates = f"{ref}A{first}:A{last}"
            hours = f"{ref}I{first}:I{last}"
            pay = f"{ref}J{first}:J{last}"
            criteria = (f'{dates},">="&DATE({ty},4,6),'
                        f'{dates},"<"&DATE({ty + 1},4,6)')
            hour_sums.append(f"SUMIFS({hours},{criteria})")
            pay_sums.append(f"SUMIFS({pay},{criteria})")
            missing_rates.append(f'COUNTIFS({criteria},{pay},"")')
        ws.cell(row=row, column=1, value=tax_year_label(ty))
        ws.cell(row=row, column=5, value="=SUM(" + ",".join(hour_sums) + ")")
        ws.cell(row=row, column=6,
                value='=IF(SUM(' + ",".join(missing_rates) + ')>0,"",SUM('
                      + ",".join(pay_sums) + '))')
        _style_summary_row(ws, row, bold=True, fill=X_MINT_SOFT)


def _style_summary_row(ws, row, bold=False, banded=False, fill=None):
    for c in range(1, 7):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=X_FONT, bold=bold, size=11, color=X_TXT)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center",
                                   vertical="center")
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        elif banded:
            cell.fill = PatternFill("solid", fgColor=X_GREY)
        if c in (3, 4, 5):
            cell.number_format = "0.0"
        elif c == 6:
            cell.number_format = "#,##0.00"
        elif c == 2:
            cell.number_format = "0"
