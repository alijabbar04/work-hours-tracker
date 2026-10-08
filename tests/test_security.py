"""Synthetic privacy and malicious-input regression tests; never use user data."""

import datetime as dt
import json
import re
import shutil
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest
import requests
from openpyxl import load_workbook

from core import ai, inbox, invoice, mobile
from core.excel_store import Entry, ExcelStore, build_workbook, read_data_rows


def synthetic_config():
    return {
        "rates": [{"effective_date": "2024-01-01", "rate": 20.0}],
        "personal": {"name": "Synthetic Worker", "email": "worker@example.invalid"},
        "bill_to": {"lines": ["Synthetic Client"]},
        "payment": {"account_name": "Synthetic Account", "sort_code": "00-00-00",
                    "account_no": "00000000"},
        "invoice": {"next_number": 99, "self_employed_note": "Synthetic note",
                    "thanks_template": "Reference {number}."},
        "defaults": {},
    }


def example_entry(**kwargs):
    values = dict(date=dt.date(2024, 2, 1), start=dt.time(9), finish=dt.time(17),
                  break_min=30, overtime=0, type="Office", summary="Example task")
    values.update(kwargs)
    return Entry(**values)


def phone_payload(**kwargs):
    values = dict(app="work-hours-tracker", version=1, date="2024-02-01", start="09:00",
                  finish="17:00", break_min=30, overtime=0, type="Office", summary="Example task")
    values.update(kwargs)
    return values


class TestExportPrivacy:
    def test_mobile_requires_explicit_opt_in(self, tmp_path):
        cfg = synthetic_config()
        cfg["onedrive_mirror_dir"] = str(tmp_path / "mirror")
        assert mobile.export_all([example_entry()], cfg) == []
        assert not (tmp_path / "mirror").exists()

    def test_mobile_omits_invoice_identity_and_payment_by_default(self):
        cfg = synthetic_config()
        payload = mobile.build_payload([example_entry()], cfg)
        html = mobile.render_html([example_entry()], cfg)
        for field in ("personal", "bill_to", "payment"):
            assert payload[field] == {}
        assert payload["invoice"] == {"next_number": 1, "self_note": "", "thanks": ""}
        for text in ("Synthetic Worker", "Synthetic Client", "00000000", "Synthetic note"):
            assert text not in html
        assert "localStorage" not in html and "fetch(" not in html

    def test_mobile_invoice_details_need_separate_opt_in(self):
        cfg = synthetic_config()
        cfg["mobile_include_invoice_details"] = True
        payload = mobile.build_payload([example_entry()], cfg)
        assert payload["personal"]["name"] == "Synthetic Worker"
        assert payload["payment"]["account_no"] == "00000000"
        assert payload["invoice"]["next_number"] == 99

    def test_json_script_roundtrip_does_not_allow_html_breakout(self):
        dangerous = '</ScRiPt><img src=x onerror="alert(1)"><!--&'
        html = mobile.render_html([example_entry(summary=dangerous, type=dangerous)], synthetic_config())
        embedded = re.search(r'<script id="whdata" type="application/json">(.*?)</script>', html, re.S).group(1)
        assert "<" not in embedded and ">" not in embedded
        assert json.loads(embedded)["entries"][0]["su"] == dangerous
        assert '${esc(e.t)}' in html
        assert '${esc(df.start??"09:00")}' in html

    @pytest.mark.skipif(shutil.which("node") is None, reason="optional Node runtime unavailable")
    def test_generated_mobile_renders_untrusted_values_safely(self, tmp_path):
        cfg = synthetic_config()
        cfg["mobile_include_invoice_details"] = True
        cfg["personal"]["name"] = '<img src=x onerror="bad()">'
        cfg["defaults"]["start"] = '" autofocus onfocus="bad()'
        out = tmp_path / "synthetic-mobile.html"
        out.write_text(mobile.render_html([
            example_entry(type='<img src=x onerror="bad()">', summary='<img src=x>')
        ], cfg), encoding="utf-8")
        subprocess.run([shutil.which("node"), str(Path(__file__).with_name("mobile_smoke.cjs")),
                        str(out)], check=True, capture_output=True, text=True, timeout=15)


class TestWorkbookSafety:
    @pytest.mark.parametrize("text", ['=HYPERLINK("https://example.invalid","click")',
                                      "+command", "@SUM(A1)", "-example"])
    def test_user_text_is_literal_after_disk_roundtrip(self, tmp_path, text):
        entry = example_entry(type=text, summary=text)
        path = tmp_path / "safe.xlsx"
        build_workbook([entry], synthetic_config()).save(path)
        wb = load_workbook(path, data_only=False)
        ws = wb["Feb 2024"]
        for column in (8, 11):
            assert ws.cell(2, column).data_type == "s"
            assert ws.cell(2, column).value == text
        assert ws.cell(2, 10).data_type == "f"  # app's own arithmetic remains a formula
        assert read_data_rows(ws)[0].summary == text
        wb.close()

    def test_logging_without_rate_keeps_hours_and_pay_unknown(self):
        cfg = synthetic_config()
        cfg["rates"] = []
        wb = build_workbook([example_entry()], cfg)
        ws = wb["Feb 2024"]
        assert ws.cell(2, 10).value is None
        assert "No rate configured" in ws.cell(2, 10).comment.text
        assert 'COUNT(J2:J2)' in ws.cell(4, 10).value
        assert 'COUNT(' in wb["Summary"].cell(4, 6).value
        assert read_data_rows(ws)[0].total_hours() == 7.5

    def test_failed_save_restores_in_memory_entries(self, tmp_path, monkeypatch):
        original = example_entry(summary="Original synthetic note")
        store = ExcelStore(str(tmp_path / "data.xlsx"), synthetic_config())
        store.entries = [original]
        def fail():
            raise PermissionError("synthetic lock")
        monkeypatch.setattr(store, "save", fail)
        with pytest.raises(PermissionError):
            store.upsert(example_entry(summary="Replacement"))
        assert store.entries == [original]
        with pytest.raises(PermissionError):
            store.delete(original.date)
        assert store.entries == [original]

    def test_april_entries_are_split_at_actual_tax_year_boundary(self):
        entries = [example_entry(date=dt.date(2024, 4, 5)),
                   example_entry(date=dt.date(2024, 4, 6))]
        wb = build_workbook(entries, synthetic_config())
        ws = wb["Summary"]
        labels = {ws.cell(row, 1).value: row for row in range(1, ws.max_row + 1)}
        assert "2023/24" in labels and "2024/25" in labels
        for label, start_year in (("2023/24", 2023), ("2024/25", 2024)):
            formula = ws.cell(labels[label], 5).value
            assert f'DATE({start_year},4,6)' in formula
            assert f'DATE({start_year + 1},4,6)' in formula
            assert "SUMIFS(" in formula
            assert "'Apr 2024'!A2:A3" in formula


class TestUntrustedPhoneFiles:
    @pytest.mark.parametrize("changes", [
        {"version": 2}, {"break_min": -1}, {"break_min": 1.5}, {"break_min": True},
        {"overtime": float("nan")}, {"overtime": float("inf")}, {"overtime": -1},
        {"overtime": 25}, {"type": "<script>bad</script>"}, {"start": "25:00"},
        {"finish": None}, {"finish": "08:00"}, {"break_min": 600},
        {"summary": ["not text"]}, {"summary": "x" * 20001},
    ])
    def test_invalid_payload_rejected_before_import(self, changes):
        with pytest.raises((ValueError, TypeError)):
            inbox.parse_entry(phone_payload(**changes))

    def test_json_list_is_reported_as_error_without_crash(self, tmp_path):
        path = tmp_path / "inbox"
        path.mkdir()
        (path / "bad.json").write_text("[]", encoding="utf-8")
        pending, errors = inbox.scan({"onedrive_mirror_dir": str(tmp_path)})
        assert not pending and len(errors) == 1

    def test_oversized_file_rejected(self, tmp_path):
        path = tmp_path / "inbox"
        path.mkdir()
        (path / "large.json").write_text(" " * (inbox.MAX_FILE_BYTES + 1), encoding="utf-8")
        pending, errors = inbox.scan({"onedrive_mirror_dir": str(tmp_path)})
        assert not pending and "256 KB" in errors[0][1]

    def test_archive_never_overwrites_previous_import(self, tmp_path):
        path = tmp_path / "entry.json"
        imported = tmp_path / "imported"
        imported.mkdir()
        (imported / path.name).write_text("old", encoding="utf-8")
        path.write_text("new", encoding="utf-8")
        assert inbox.archive(str(path)) is True
        assert sorted(p.read_text() for p in imported.iterdir()) == ["new", "old"]

    def test_locked_phone_import_restores_memory_and_leaves_source(self, monkeypatch):
        from ui import phone_sync
        original = example_entry(summary="Original synthetic note")
        replacement = example_entry(summary="Replacement synthetic note")
        def fail_save():
            raise PermissionError("synthetic lock")
        store = SimpleNamespace(entries=[original], entry_for_date=lambda date: original,
                                save=fail_save)
        app = SimpleNamespace(cfg={}, store=store,
                              on_data_changed=lambda: pytest.fail("save failed"))
        notices = []
        monkeypatch.setattr(phone_sync.inbox, "scan", lambda cfg: ([("synthetic.json", replacement)], []))
        monkeypatch.setattr(phone_sync.inbox, "archive", lambda path: pytest.fail("must not archive"))
        monkeypatch.setattr(phone_sync.messagebox, "askyesno", lambda *args, **kwargs: True)
        monkeypatch.setattr(phone_sync.messagebox, "showerror", lambda *args, **kwargs: notices.append(args))
        assert phone_sync.import_phone_entries(None, app) == 0
        assert store.entries == [original] and notices

    def test_duplicate_phone_dates_never_overwrite_by_file_order(self, monkeypatch):
        from ui import phone_sync
        first, second = example_entry(summary="First"), example_entry(summary="Second")
        store = SimpleNamespace(entries=[], entry_for_date=lambda date: None,
                                save=lambda: pytest.fail("ambiguous import must not save"))
        app = SimpleNamespace(cfg={}, store=store)
        monkeypatch.setattr(phone_sync.inbox, "scan", lambda cfg: (
            [("first.json", first), ("second.json", second)], []))
        monkeypatch.setattr(phone_sync.messagebox, "askyesno", lambda *args, **kwargs: True)
        warnings = []
        monkeypatch.setattr(phone_sync.messagebox, "showwarning", lambda *args, **kwargs: warnings.append(args))
        assert phone_sync.import_phone_entries(None, app) == 0
        assert store.entries == [] and warnings


class TestInvoiceSafety:
    def test_filename_cannot_escape_invoice_folder(self):
        cfg = synthetic_config()
        cfg["personal"]["name"] = "../outside\\folder:<bad>"
        name = invoice.default_filename(cfg, 2024, 2)
        assert "/" not in name and "\\" not in name and ":" not in name
        assert name.endswith("_February_2024.pdf")

    def test_pdf_treats_untrusted_markup_as_text(self, tmp_path):
        cfg = synthetic_config()
        cfg["personal"]["name"] = '<img src="https://example.invalid/test"/> & Worker'
        cfg["bill_to"]["lines"] = ["Client <unclosed-tag> & example"]
        cfg["payment"]["account_name"] = "A <b>literal</b> name"
        cfg["invoice"]["thanks_template"] = "Ref {number}; keep {unknown} literal."
        data = invoice.build_invoice_data([example_entry(type="<b>literal</b> & work")], cfg, 2024, 2, 1)
        out = tmp_path / "safe.pdf"
        invoice.generate_pdf(str(out), data, cfg)
        assert out.stat().st_size > 1000
        assert invoice._text('<img src="example"/>') == '&lt;img src="example"/&gt;'


class TestAISecrets:
    def test_plaintext_keyring_refused(self, monkeypatch):
        unsafe = type("PlaintextStore", (), {"__module__": "keyrings.alt.file", "priority": 1})()
        fake = SimpleNamespace(get_keyring=lambda: unsafe,
                               set_password=lambda *args: pytest.fail("must not store key"))
        monkeypatch.setattr(ai, "keyring", fake)
        with pytest.raises(RuntimeError, match="Secure OS"):
            ai.save_api_key("synthetic-token")

    def test_secure_os_keyring_stores_namespaced_key(self, monkeypatch):
        backend = type("WinVault", (), {"__module__": "keyring.backends.Windows", "priority": 5})()
        calls = []
        fake = SimpleNamespace(get_keyring=lambda: backend,
                               set_password=lambda *args: calls.append(args))
        monkeypatch.setattr(ai, "keyring", fake)
        ai.save_api_key(" synthetic-token ")
        assert calls == [(ai.KEYRING_SERVICE, ai.KEYRING_USER, "synthetic-token")]
        assert "WorkHoursTracker" in ai.KEYRING_SERVICE

    def test_network_error_never_exposes_notes_or_token(self, monkeypatch):
        def fail(*args, **kwargs):
            raise requests.ConnectionError("synthetic-token private-synthetic-notes")
        monkeypatch.setattr(ai.requests, "post", fail)
        with pytest.raises(RuntimeError) as err:
            ai.claude_bullets("synthetic-token", "private-synthetic-notes")
        assert "synthetic-token" not in str(err.value)
        assert "private-synthetic-notes" not in str(err.value)

    def test_api_error_body_not_displayed(self, monkeypatch):
        monkeypatch.setattr(ai.requests, "post", lambda *args, **kwargs: SimpleNamespace(
            ok=False, status_code=401, text="synthetic-token private-synthetic-notes"))
        with pytest.raises(RuntimeError, match="HTTP 401") as err:
            ai.claude_bullets("synthetic-token", "private-synthetic-notes")
        assert "synthetic-token" not in str(err.value)
