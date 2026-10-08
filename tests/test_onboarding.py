"""Setup/import regression tests with arbitrary synthetic data and temporary files."""

import copy
import datetime as dt
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core import config as cfgmod, migrate, reminder
from core.excel_store import Entry, build_workbook
from ui.setup_view import initial_rate


@pytest.fixture
def isolated_config(tmp_path, monkeypatch):
    data = tmp_path / "local-data"
    monkeypatch.setattr(cfgmod, "DATA_DIR", str(data))
    monkeypatch.setattr(cfgmod, "BACKUP_DIR", str(data / "backups"))
    monkeypatch.setattr(cfgmod, "INVOICE_DIR", str(data / "invoices"))
    monkeypatch.setattr(cfgmod, "WORKBOOK_PATH", str(data / "Work_Hours_Tracker.xlsx"))
    monkeypatch.setattr(cfgmod, "CONFIG_PATH", str(data / "config.json"))
    return copy.deepcopy(cfgmod.DEFAULT_CONFIG)


def sample_entries():
    return [Entry(dt.date(2024, 3, 8), dt.time(9), dt.time(17), 30, 1, "Office", "Synthetic task A"),
            Entry(dt.date(2024, 3, 11), dt.time(10), dt.time(18), 45, 0, "Office", "Synthetic task B")]


def write_source(path, cfg, legacy=False):
    workbook = build_workbook(sample_entries(), cfg)
    if legacy:
        workbook.remove(workbook["Summary"])
        workbook["Mar 2024"].title = "Timesheet"
    workbook.save(path)
    workbook.close()


def test_new_install_is_blank_and_local_only(isolated_config):
    cfg = cfgmod.load_config()
    assert cfg["rates"] == []
    assert cfg["personal"] == {"name": "", "address_lines": [], "email": "", "phone": ""}
    assert cfg["payment"] == {"account_name": "", "sort_code": "", "account_no": ""}
    assert cfg["bill_to"]["lines"] == []
    assert cfg["invoice"]["self_employed_note"] == ""
    assert cfg["onedrive_mirror_dir"] == ""
    for flag in ("mobile_export_enabled", "mobile_include_invoice_details", "mirror_invoices",
                 "reminder_enabled", "ai_notes_consent", "onboarding_complete"):
        assert cfg[flag] is False
    assert not Path(cfgmod.CONFIG_PATH).exists()
    assert not Path(cfgmod.WORKBOOK_PATH).exists()


def test_config_roundtrip_and_bad_json_are_not_silently_overwritten(isolated_config):
    isolated_config["rates"] = [{"effective_date": "2024-01-01", "rate": 21.75}]
    cfgmod.save_config(isolated_config)
    assert cfgmod.load_config()["rates"] == isolated_config["rates"]
    path = Path(cfgmod.CONFIG_PATH)
    path.write_text("{invalid", encoding="utf-8")
    with pytest.raises(json.JSONDecodeError):
        cfgmod.load_config()
    assert path.read_text(encoding="utf-8") == "{invalid"


@pytest.mark.parametrize("bad", [[], {"personal": []}, {"rates": "unexpected"}])
def test_config_rejects_wrong_structures(isolated_config, bad):
    cfgmod.ensure_dirs()
    Path(cfgmod.CONFIG_PATH).write_text(json.dumps(bad), encoding="utf-8")
    with pytest.raises(ValueError):
        cfgmod.load_config()


def test_explicit_data_directory_override(tmp_path, monkeypatch):
    monkeypatch.setenv("WORK_HOURS_TRACKER_DATA_DIR", str(tmp_path / "sandbox"))
    assert cfgmod.data_directory() == str(tmp_path / "sandbox")


def test_optional_initial_rate():
    assert initial_rate("", "") is None
    assert initial_rate("0", "2024-01-01") == {"effective_date": "2024-01-01", "rate": 0.0}
    assert initial_rate("21.75", "2024-01-01")["rate"] == 21.75


@pytest.mark.parametrize("rate", ["-1", "nan", "inf", "text"])
def test_initial_rate_rejects_invalid_numbers(rate):
    with pytest.raises(ValueError):
        initial_rate(rate, "2024-01-01")


def test_initial_rate_requires_valid_date():
    with pytest.raises(ValueError):
        initial_rate("21.75", "not-a-date")


@pytest.mark.parametrize("legacy", [False, True])
def test_import_preserves_source_and_backs_up_previous_local_data(isolated_config, tmp_path, legacy):
    isolated_config["rates"] = [{"effective_date": "2024-01-01", "rate": 21.75}]
    source = tmp_path / "selected.xlsx"
    write_source(source, isolated_config, legacy)
    source_bytes = source.read_bytes()
    cfgmod.ensure_dirs()
    old = build_workbook([Entry(dt.date(2023, 12, 1), overtime=3.0)], isolated_config)
    old.save(cfgmod.WORKBOOK_PATH)
    old.close()
    old_bytes = Path(cfgmod.WORKBOOK_PATH).read_bytes()
    ok, report = migrate.run(isolated_config, {}, str(source))
    assert ok, report
    assert source.read_bytes() == source_bytes
    backups = list(Path(cfgmod.BACKUP_DIR).glob("*.xlsx"))
    assert len(backups) == 2
    assert source_bytes in [path.read_bytes() for path in backups]
    assert old_bytes in [path.read_bytes() for path in backups]
    migrated = migrate.read_original(cfgmod.WORKBOOK_PATH)
    assert migrated.entries == sample_entries()
    assert any("computed pay agrees" in line for line in report)


def test_import_missing_rates_still_preserves_hours(isolated_config, tmp_path):
    source = tmp_path / "selected.xlsx"
    write_source(source, isolated_config)
    ok, report = migrate.run(isolated_config, {}, str(source))
    assert ok, report
    assert migrate.read_original(cfgmod.WORKBOOK_PATH).entries == sample_entries()
    assert any("Pay not calculated" in line for line in report)


def test_dateless_import_requires_confirmation_and_rejects_duplicates(isolated_config, tmp_path):
    source = tmp_path / "selected.xlsx"
    write_source(source, isolated_config, legacy=True)
    wb = load_workbook(source)
    wb["Timesheet"]["A3"] = None
    wb.save(source)
    wb.close()
    source_bytes = source.read_bytes()
    plan = migrate.read_original(source)
    assert plan.dateless[0].row == 3
    assert plan.dateless[0].suggested is None
    assert not migrate.run(isolated_config, {}, str(source))[0]
    assert not Path(cfgmod.WORKBOOK_PATH).exists()
    assert not migrate.run(isolated_config, {3: dt.date(2024, 3, 8)}, str(source))[0]
    assert source.read_bytes() == source_bytes
    ok, report = migrate.run(isolated_config, {3: dt.date(2024, 3, 11)}, str(source))
    assert ok, report
    assert migrate.read_original(cfgmod.WORKBOOK_PATH).entries == sample_entries()


def test_failed_verification_does_not_replace_local_or_source(isolated_config, tmp_path, monkeypatch):
    source = tmp_path / "selected.xlsx"
    write_source(source, isolated_config)
    cfgmod.ensure_dirs()
    Path(cfgmod.WORKBOOK_PATH).write_bytes(b"existing local data")
    source_bytes = source.read_bytes()
    monkeypatch.setattr(migrate, "verify", lambda *_: (False, ["synthetic verification failure"]))
    assert not migrate.run(isolated_config, {}, str(source))[0]
    assert Path(cfgmod.WORKBOOK_PATH).read_bytes() == b"existing local data"
    assert source.read_bytes() == source_bytes
    assert not list(Path(cfgmod.DATA_DIR).glob("*.import.xlsx"))


def test_verification_detects_modified_values_and_pay_formula(isolated_config, tmp_path):
    isolated_config["rates"] = [{"effective_date": "2024-01-01", "rate": 21.75}]
    source = tmp_path / "selected.xlsx"
    write_source(source, isolated_config)
    wb = load_workbook(source)
    wb["Mar 2024"]["J2"] = "=I2*999"
    wb["Mar 2024"]["K3"] = "Changed task"
    wb.save(source)
    wb.close()
    ok, report = migrate.verify(sample_entries(), str(source), isolated_config)
    assert not ok
    assert any("raw entry values" in line for line in report)
    assert any("pay formula" in line for line in report)


def test_import_refuses_source_equal_to_target(isolated_config, tmp_path):
    source = tmp_path / "selected.xlsx"
    write_source(source, isolated_config)
    original = source.read_bytes()
    with pytest.raises(ValueError, match="already the local"):
        migrate.run(isolated_config, {}, str(source), str(source))
    assert source.read_bytes() == original


@pytest.mark.parametrize("cell,value", [("C2", "unknown time"), ("E2", 2.5), ("G2", -1), ("A2", "not a date")])
def test_import_rejects_ambiguous_fields_without_replacing_local(isolated_config, tmp_path, cell, value):
    source = tmp_path / "selected.xlsx"
    write_source(source, isolated_config)
    wb = load_workbook(source)
    wb["Mar 2024"][cell] = value
    wb.save(source)
    wb.close()
    with pytest.raises(ValueError):
        migrate.run(isolated_config, {}, str(source))
    assert not Path(cfgmod.WORKBOOK_PATH).exists()


def test_reminders_are_portable_and_disabled_for_custom_data_path(monkeypatch):
    monkeypatch.setattr(reminder.sys, "platform", "linux")
    assert reminder.is_enabled() is False
    reminder.disable()
    with pytest.raises(RuntimeError, match="Windows only"):
        reminder.enable()
    monkeypatch.setattr(reminder.sys, "platform", "win32")
    monkeypatch.setenv("WORK_HOURS_TRACKER_DATA_DIR", "sandbox")
    with pytest.raises(RuntimeError, match="custom data"):
        reminder.enable()
